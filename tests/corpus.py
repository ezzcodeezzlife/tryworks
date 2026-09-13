"""Build the test and compatibility corpus: one realistic document per supported format.

Office files are written with python-docx, python-pptx and openpyxl (development dependencies),
so they have the XML structure real tools produce. The PDF is assembled by hand with standard
Helvetica fonts. Everything is generated, so the corpus carries no third-party content.

    python tests/corpus.py OUTPUT_DIR
"""

from __future__ import annotations

import datetime as dt
import sys
from email.message import EmailMessage
from pathlib import Path

COMPANY = "Keel & Rope Supply Co."


# -- DOCX ------------------------------------------------------------------------------------


def _docx_hyperlink(paragraph, text: str, url: str) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_docx(path: Path) -> Path:
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = f"{COMPANY} | Board memo"
    section.footer.paragraphs[0].text = "Internal distribution only"

    doc.add_heading("Quarterly Operations Memo", level=0)
    doc.add_heading("Summary", level=1)
    p = doc.add_paragraph()
    p.add_run("Revenue grew ")
    p.add_run("eleven percent").bold = True
    p.add_run(" year over year, driven by ")
    p.add_run("marine hardware").italic = True
    p.add_run(" sales across the northern ports.")

    p = doc.add_paragraph("The full figures are in the ")
    _docx_hyperlink(p, "third quarter report", "https://example.com/q3-report")
    p.add_run(" that finance published last week.")

    doc.add_heading("Priorities", level=2)
    doc.add_paragraph("Expand the Rotterdam warehouse", style="List Bullet")
    doc.add_paragraph("Hire two procurement analysts", style="List Bullet")
    doc.add_paragraph("Close the third quarter audit", style="List Number")

    table = doc.add_table(rows=3, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Warehouse totals"
    table.cell(0, 2).text = "Units"
    for r, (name, sku, units) in enumerate([("Rotterdam", "KR-1001", "1200"), ("Antwerp", "KR-2040", "860")], start=1):
        table.cell(r, 0).text = name
        table.cell(r, 1).text = sku
        table.cell(r, 2).text = units

    # -- Word records where pages broke when it last laid the document out --
    p = doc.add_paragraph()
    p.add_run()._r.append(OxmlElement("w:lastRenderedPageBreak"))
    p.add_run("The outlook for the fourth quarter remains positive despite higher fuel costs.")
    doc.add_heading("Contacts", level=1)
    doc.add_paragraph("operations@example.com")
    doc.save(path)
    return path


# -- PPTX ------------------------------------------------------------------------------------


def _pptx_bullet(paragraph, char: str = "•") -> None:
    from lxml import etree
    from pptx.oxml.ns import qn

    ppr = paragraph._p.get_or_add_pPr()
    etree.SubElement(ppr, qn("a:buChar"), char=char)


def make_pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Fleet Maintenance Review"
    slide.placeholders[1].text = "Prepared by the operations team for the September planning cycle"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Priorities"
    frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame
    frame.text = "Replace worn rigging on three vessels"
    _pptx_bullet(frame.paragraphs[0])
    para = frame.add_paragraph()
    para.text = "Schedule dry dock inspections"
    para.level = 1
    _pptx_bullet(para)
    note = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(8), Inches(1)).text_frame
    note.text = "The inspection backlog has grown since March and now affects twelve vessels."

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Budget"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate([("Item", "Cost"), ("Rigging", "$42,000"), ("Dry dock", "$118,500")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    slide.notes_slide.notes_text_frame.text = "Mention that the dry dock quote expires in October."
    prs.save(path)
    return path


# -- XLSX ------------------------------------------------------------------------------------


def make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"
    ws["A1"] = "Inventory by warehouse"
    ws.append([])
    ws.append(["Warehouse", "SKU", "Units"])
    ws.append(["Rotterdam", "KR-1001", 1200])
    ws.append(["Antwerp", "KR-2040", 860])
    ws.append(["Hamburg", "KR-3300", 415])
    ws.append([])
    ws.append(["Counts reflect the August cycle count."])

    ws2 = wb.create_sheet("Shipments")
    ws2.append(["Date", "Vessel", "Tons"])
    ws2.append([dt.date(2026, 8, 1), "Marlin", 42.5])
    ws2.append([dt.date(2026, 8, 3), "Petrel", 38])
    wb.save(path)
    return path


# -- PDF -------------------------------------------------------------------------------------


def _pdf_text(value: str) -> bytes:
    raw = value.encode("cp1252")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def build_pdf(pages: list[list[tuple[str, int, float, float, str]]], width: int = 612, height: int = 792) -> bytes:
    """Assemble a PDF. Each page is a list of (font, size, x, y, text); font is "F1" or "F2" (bold)."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    regular = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    pages_ref = add(b"")
    kids = []
    for items in pages:
        content = b"".join(
            b"BT /%s %d Tf %.2f %.2f Td (%s) Tj ET\n" % (font.encode(), size, x, y, _pdf_text(text))
            for font, size, x, y, text in items
        )
        stream = add(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
        kids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 %d %d] /Resources << /Font << /F1 %d 0 R /F2 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_ref, width, height, regular, bold, stream)
            )
        )
    objects[pages_ref - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % k for k in kids),
        len(kids),
    )
    catalog = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_ref)

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    return bytes(out)


def make_pdf(path: Path) -> Path:
    page1 = [
        ("F1", 9, 72, 770, f"{COMPANY} | Operations Report"),
        ("F2", 20, 72, 715, "Operations Report"),
        ("F1", 11, 72, 685, "The third quarter closed with record throughput at the Rotterdam"),
        ("F1", 11, 72, 671, "warehouse. Order volume rose eleven percent while fulfilment times"),
        ("F1", 11, 72, 657, "fell to under two days for most customers."),
        ("F1", 11, 72, 625, "• Rotterdam throughput up eleven percent"),
        ("F1", 11, 72, 611, "• Two new procurement analysts hired"),
        ("F2", 14, 72, 575, "Next Steps"),
        ("F1", 11, 72, 553, "The board approved funding for a second warehouse in Antwerp."),
        ("F1", 11, 72, 539, "Construction is expected to begin in the spring."),
        ("F1", 9, 72, 30, "Page 1 of 2"),
    ]
    page2 = [
        ("F2", 16, 72, 715, "Regional Detail"),
        ("F1", 11, 72, 685, "Northern revenue grew steadily through"),
        ("F1", 11, 72, 671, "the quarter as contracts with ferry"),
        ("F1", 11, 72, 657, "operators came online in July."),
        ("F1", 11, 330, 685, "Southern margins improved after the team"),
        ("F1", 11, 330, 671, "renegotiated fuel surcharges with three"),
        ("F1", 11, 330, 657, "of its largest shipping partners."),
        ("F1", 9, 72, 30, "Page 2 of 2"),
    ]
    path.write_bytes(build_pdf([page1, page2]))
    return path


# -- text formats ----------------------------------------------------------------------------

HTML = """<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><title>Harbor Notes</title>
<style>body { font-family: sans-serif; }</style><script>track();</script></head>
<body>
<nav><a href="/">Home</a> | <a href="/about">About</a></nav>
<h1>Harbor Operations Update</h1>
<p>The harbor authority <b>extended</b> night shifts after a surge in container traffic.
Crews now work until <i>two in the morning</i> on weekdays.</p>
<h2>What changed</h2>
<ul>
  <li>Night shifts run five days a week</li>
  <li>Two cranes were added<ul><li>One arrives in October</li></ul></li>
</ul>
<table>
  <thead><tr><th>Terminal</th><th>Moves per hour</th></tr></thead>
  <tbody><tr><td>North</td><td>31</td></tr><tr><td>South</td><td>27</td></tr></tbody>
</table>
<pre>berth_schedule --week 37
  north: 14 vessels</pre>
<p>Read the <a href="https://example.com/notice">full notice</a> for shift details.</p>
<img src="https://example.com/cranes.png" alt="New cranes at the north terminal">
<footer>Published by the harbor authority press office.</footer>
</body></html>
"""

MARKDOWN = """# Onboarding Guide

Welcome to the procurement team. This guide explains how purchase orders move through approval.

## Approval steps

1. Submit the order in the purchasing portal
2. Your manager approves orders under ten thousand euros
3. Finance reviews anything larger

- Keep receipts for every order
- Tag urgent orders with **priority**

| Amount | Approver |
|--------|----------|
| Under 10k | Manager |
| 10k and over | Finance |

```bash
po submit --vendor 42 --amount 9800
```

Questions go to [the procurement desk](https://example.com/desk).
"""

TEXT = """SHIPPING POLICY

Orders placed before noon ship the same business day. Orders placed after noon ship the next business day.

• Standard delivery takes three to five days
• Express delivery takes one to two days

Returns are accepted within thirty days of delivery if the item is unused and in its original packaging.

support@example.com

Doylestown, PA 18901
"""

CSV = """Vessel,Port,Arrival,Tons
Marlin,Rotterdam,2026-08-01,42.5
Petrel,Antwerp,2026-08-03,38
"Albatross, II",Hamburg,2026-08-05,51
"""


def make_eml(path: Path) -> Path:
    msg = EmailMessage()
    msg["From"] = "Dana Whitfield <dana@example.com>"
    msg["To"] = "ops@example.com, Lee Park <lee@example.com>"
    msg["Cc"] = "finance@example.com"
    msg["Subject"] = "Dry dock schedule for October"
    msg["Date"] = "Tue, 08 Sep 2026 09:30:00 +0200"
    msg["Message-ID"] = "<schedule-1234@example.com>"
    msg.set_content(
        "Hi team,\n\nThe dry dock slot for the Marlin moved to the second week of October.\n\nThanks,\nDana\n"
    )
    msg.add_alternative(
        "<html><body><p>Hi team,</p><p>The dry dock slot for the <b>Marlin</b> moved to the second week of October.</p>"
        "<p>Thanks,<br>Dana</p></body></html>",
        subtype="html",
    )
    msg.add_attachment(
        b"Dry dock checklist\n\nDrain the ballast tanks before the vessel enters the dock.\n",
        maintype="text",
        subtype="plain",
        filename="checklist.txt",
    )
    path.write_bytes(msg.as_bytes())
    return path


def build_corpus(out_dir: Path) -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "html": out_dir / "harbor.html",
        "md": out_dir / "onboarding.md",
        "txt": out_dir / "policy.txt",
        "csv": out_dir / "arrivals.csv",
    }
    files["html"].write_text(HTML, encoding="utf-8", newline="\n")
    files["md"].write_text(MARKDOWN, encoding="utf-8", newline="\n")
    files["txt"].write_text(TEXT, encoding="utf-8", newline="\n")
    files["csv"].write_text(CSV, encoding="utf-8", newline="\n")
    files["eml"] = make_eml(out_dir / "schedule.eml")
    files["pdf"] = make_pdf(out_dir / "report.pdf")
    files["docx"] = make_docx(out_dir / "memo.docx")
    files["pptx"] = make_pptx(out_dir / "fleet.pptx")
    files["xlsx"] = make_xlsx(out_dir / "inventory.xlsx")
    return files


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "corpus")
    for kind, file in build_corpus(target).items():
        print(f"{kind:5} {file}")
